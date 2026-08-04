"""endpoint ของคลัง skills

⚠️ กติกา 2 ข้อของไฟล์นี้:

1. **`async def` ห้ามเรียกของ sync ตรงๆ** — handler ที่ต้องอ่าน body ถูกบังคับให้เป็น
   `async` งาน sync ที่ตามมา (LLM / เขียนไฟล์บน volume NAS / parse เอกสาร) จึงรันบน
   event loop = ทุกคำขอของทุกคนหยุดรอ → ต้องผ่าน `run_in_threadpool`
   handler ที่เป็น `def` ธรรมดา FastAPI โยนเข้า threadpool ให้เองอยู่แล้ว

2. **ห้าม `_load_skills_db()` → แก้ → `_save_skills_db()` เองในไฟล์นี้** — ข้อ 1 ย้ายโค้ด
   ที่เคยรันทีละอันไปรันพร้อมกัน read-modify-write ที่ไม่มี lock จึงกลายเป็น lost update
   ทันที ให้ผ่าน `set_skill_entry()` / `delete_skill_entries()` ซึ่งถือ `_db_lock`
   (ดู tests/test_skills_db_concurrency.py + tests/test_memory_skills_router_concurrency.py)
"""
import base64
import io
import json
import logging
import os
from datetime import datetime
from fastapi import APIRouter, Request, UploadFile, File
from starlette.concurrency import run_in_threadpool

from core.config import SKILLS_DIR
from utils.skills import (
    get_skill_count, auto_extract_skills, _load_skills_db,
    cleanup_junk_skills, set_skill_entry, delete_skill_entries, SkillsDbError,
)
from utils.llm import stream_response

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api", tags=["skills"])


def _collect_stream(msgs) -> str:
    """หมุน sync generator ของ LLM ให้จบ — ต้องรันในเธรด ไม่ใช่บน event loop

    `chat.py` เรียก `stream_response()` ตัวเดียวกันแล้วปลอดภัย เพราะ**ส่งต่อ** generator
    ให้ `StreamingResponse` ซึ่ง starlette ห่อด้วย `iterate_in_threadpool()` ให้เอง
    ที่นี่เรา `"".join()` เอง → ถ้าไม่ย้ายเข้า threadpool = Gemini call เต็มรอบบน event loop
    "sync generator ปลอดภัย" เป็นจริงเฉพาะตอนที่ starlette เป็นคนหมุนเท่านั้น
    """
    return "".join(stream_response(msgs, provider="gemini"))


def _write_skill_file(path: str, text: str) -> None:
    """เขียน .md ลง SKILLS_DIR — บน prod เป็น volume ที่ mount จาก NAS จึงไม่ใช่ดิสก์ local"""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


@router.get("/skills")
def list_skills():
    db = _load_skills_db()
    skills_dir = SKILLS_DIR
    md_files = [f for f in os.listdir(skills_dir) if f.endswith(".md")] if os.path.isdir(skills_dir) else []
    chroma_count = get_skill_count()
    return {"skills": db, "count": len(db), "md_files": len(md_files), "chroma_count": chroma_count}


@router.get("/skills/list")
def skills_list():
    skills_dir = SKILLS_DIR
    if not os.path.isdir(skills_dir):
        return {"files": []}
    files = [
        {"name": f, "size": os.path.getsize(os.path.join(skills_dir, f))}
        for f in sorted(os.listdir(skills_dir))
        if os.path.isfile(os.path.join(skills_dir, f)) and f.endswith(".md")
    ]
    return {"files": files}


@router.post("/skills/extract")
async def skills_extract(request: Request):
    data = await request.json()
    content = data.get("content", "").strip()
    topic = data.get("topic", "").strip()
    if not content:
        return {"ok": False, "error": "ไม่มี content"}
    if not topic:
        topic = f"skill-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
    safe_topic = "".join(c if c.isalnum() or c in "-_" else "-" for c in topic.lower()).strip("-")
    filename = f"{safe_topic}.md"
    filepath = os.path.join(SKILLS_DIR, filename)

    msgs = [
        {"role": "system", "content": (
            "คุณคือ Technical Writer ที่เชี่ยวชาญ\n"
            "งาน: อ่าน content ด้านล่าง แล้วสกัดออกมาเป็น Skill Reference .md ที่ดี\n"
            "รูปแบบที่ต้องการ:\n"
            "- ใช้ # สำหรับชื่อหัวข้อหลัก\n"
            "- ใช้ ## สำหรับแต่ละ subtopic\n"
            "- ใส่ code block ``` ทุกครั้งที่มี code\n"
            "- สรุปกระชับ อ่านง่าย เป็น quick reference\n"
            "- ตอบเป็น markdown ล้วนๆ ไม่ต้องมีคำอธิบายเพิ่ม\n"
            f"- ชื่อหัวข้อหลัก: {topic}"
        )},
        {"role": "user", "content": content[:6000]},
    ]
    try:
        md_content = await run_in_threadpool(_collect_stream, msgs)
    except Exception as e:
        return {"ok": False, "error": f"Gemini error: {e}"}

    # เกณฑ์เดียวกับตอนลบ — กันผลลัพธ์ที่ LLM คืนมาเป็นข้อความ error/ตอบรับสั้นๆ
    # กลายเป็นไฟล์ skill (ที่มาของ `ได-เลย.md` ที่ต้องมาไล่ลบทีหลังในข้อ 9)
    from utils.skills import _is_meaningful_skill
    if not _is_meaningful_skill(topic, md_content):
        logger.info(f"[skills_extract] ปฏิเสธผลลัพธ์ที่ไม่ผ่านเกณฑ์: {topic!r}")
        return {"ok": False, "error": f"ผลลัพธ์ไม่ผ่านเกณฑ์คุณภาพ skill: {topic!r}"}

    try:
        await run_in_threadpool(_write_skill_file, filepath, md_content)
    except Exception as e:
        return {"ok": False, "error": f"บันทึกไฟล์ไม่ได้: {e}"}

    # ไฟล์ .md เขียนไปแล้ว ณ จุดนี้ — ถ้า db เขียนไม่ลงคือ **สำเร็จครึ่งเดียว**
    # ไม่กลืนเงียบ: บอกผู้เรียกด้วย `db_updated` + log ERROR (เดิม log warning แล้วตอบ
    # ok:True เฉยๆ = ผู้ใช้เห็นว่าสร้าง skill สำเร็จทั้งที่ค้นหาไม่เจอเพราะไม่มีใน db)
    db_updated = True
    warning = None
    try:
        await run_in_threadpool(set_skill_entry, topic, {
            "summary": md_content[:300].strip(),
            "source": filename,
            "updated": datetime.now().isoformat(),
        })
    except Exception as e:
        db_updated = False
        warning = f"เขียนไฟล์ .md สำเร็จ แต่บันทึกลง skills_db ไม่ได้: {e}"
        logger.error(f"[skills_extract] {warning}")

    resp = {"ok": True, "filename": filename, "path": filepath,
            "preview": md_content[:500], "db_updated": db_updated}
    if warning:
        resp["warning"] = warning
    return resp


@router.delete("/skills/{skill_id}")
def skills_delete(skill_id: str, delete_file: bool = False):
    """ลบ skill entry จาก skills_db.json (+ optional .md file)

    Args:
        skill_id: topic หรือ filename
        delete_file: ถ้า true ลบ .md file ด้วย (default false — safe)
                     SAFETY: ก่อนหน้านี้ default ลบ .md เสมอ — เปลี่ยนเป็น opt-in
                     กัน data loss จาก cleanup
    """
    from urllib.parse import unquote
    skill_id = unquote(skill_id)
    deleted_file = False
    skills_dir = SKILLS_DIR

    if delete_file:
        if ".." not in skill_id and "/" not in skill_id:
            for fname in [skill_id, f"{skill_id}.md"]:
                fp = os.path.join(skills_dir, fname)
                if os.path.exists(fp):
                    os.remove(fp)
                    deleted_file = True
                    logger.warning(f"[skills_delete] removed file {fp} (delete_file=true)")
                    break

    try:
        deleted_db = delete_skill_entries([skill_id, skill_id.replace(".md", "")])
    except SkillsDbError as e:
        logger.error(f"[skills_delete] ลบ {skill_id!r} ไม่ได้: {e}")
        return {"ok": False, "error": str(e), "deleted_file": deleted_file}

    if deleted_file or deleted_db:
        return {"ok": True, "deleted_file": deleted_file, "deleted_db": deleted_db}
    return {"ok": False, "error": "ไม่พบ skill นี้"}


@router.get("/skills/discover")
def skills_discover(days: int = 30, min_cluster: int = 3, threshold: float = 0.72):
    """สแกน chat history หา topic ที่ user ถามบ่อย → return proposals"""
    from utils.skill_discovery import discover_skills
    proposals = discover_skills(
        days=days, min_cluster=min_cluster, threshold=threshold,
    )
    return {"proposals": [p.to_dict() for p in proposals], "count": len(proposals)}


@router.post("/skills/discover/accept")
async def skills_discover_accept(request: Request):
    """รับ proposal_id → สร้าง .md ใน skills/"""
    from utils.skill_discovery import accept_proposal
    data = await request.json()
    pid = (data.get("proposal_id") or "").strip()
    if not pid:
        return {"ok": False, "error": "proposal_id required"}
    # เขียนไฟล์ .md + skills_db + sync ChromaDB — sync ทั้งเส้น
    return await run_in_threadpool(
        accept_proposal,
        pid,
        custom_topic=data.get("topic"),
        custom_content=data.get("content"),
    )


@router.get("/skills/discover/cached")
def skills_discover_cached():
    """proposals ที่รอ accept อยู่ใน memory"""
    from utils.skill_discovery import list_cached_proposals
    return {"proposals": list_cached_proposals()}


@router.post("/admin/cleanup-skills")
def cleanup_skills_endpoint():
    """ลบ junk skills ออกจาก db + re-sync ChromaDB"""
    try:
        result = cleanup_junk_skills()
    except SkillsDbError as e:
        logger.error(f"[cleanup_skills] {e}")
        return {"ok": False, "error": str(e)}
    return {"ok": True, **result}


@router.post("/admin/sync-skills")
def sync_skills_endpoint():
    try:
        from utils.skills_search import sync_skills_to_search
        db = _load_skills_db()
        if not db:
            return {"ok": False, "error": "No skills found in skills_db.json"}
        sync_skills_to_search(db)
        return {"ok": True, "synced": len(db), "message": f"Synced {len(db)} skills to ChromaDB"}
    except Exception as e:
        return {"ok": False, "error": str(e)}


def _parse_upload(content: bytes, name: str, mime: str) -> dict:
    """แกะไฟล์ที่อัปโหลด → dict ที่ handler คืนตรงๆ

    ทั้งก้อนเป็น sync และหนักจริง: pypdf/python-docx/openpyxl parse ทั้งไฟล์ (CPU)
    แล้วต่อด้วย `auto_extract_skills()` ซึ่งเขียน ChromaDB ต่อ topic ที่สกัดได้
    → เรียกจาก handler ผ่าน `run_in_threadpool` เท่านั้น
    """
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""

    if mime.startswith("image/") or ext in ('jpg', 'jpeg', 'png', 'gif', 'webp', 'bmp'):
        b64 = base64.b64encode(content).decode()
        return {"ok": True, "filename": name, "is_image": True, "b64": b64, "mime": mime or "image/jpeg"}

    if ext == "pdf" or mime == "application/pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages_text = [page.extract_text() or "" for page in reader.pages]
            raw_text = "\n\n".join(pages_text)
            text = f"[PDF: {name} — {len(reader.pages)} หน้า]\n{raw_text}"
        except ImportError:
            return {"ok": False, "error": "ไม่พบ library pypdf"}
        except Exception as e:
            return {"ok": False, "error": f"อ่าน PDF ไม่ได้: {e}"}
    elif ext == "docx" or "wordprocessingml" in mime:
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            raw_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            text = f"[DOCX: {name}]\n{raw_text}"
        except ImportError:
            return {"ok": False, "error": "ไม่พบ library python-docx"}
        except Exception as e:
            return {"ok": False, "error": f"อ่าน DOCX ไม่ได้: {e}"}
    elif ext in ("xlsx", "xls") or "spreadsheetml" in mime:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts = []
            for sheet in wb.worksheets:
                rows = ["\t".join(str(c) if c is not None else "" for c in row)
                        for row in sheet.iter_rows(values_only=True)
                        if any(c is not None for c in row)]
                if rows:
                    parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
            raw_text = "\n\n".join(parts)
            text = f"[Excel: {name}]\n{raw_text}"
        except ImportError:
            return {"ok": False, "error": "ไม่พบ library openpyxl"}
        except Exception as e:
            return {"ok": False, "error": f"อ่าน Excel ไม่ได้: {e}"}
    else:
        try:
            if ext == "json":
                parsed = json.loads(content)
                raw_text = json.dumps(parsed, ensure_ascii=False, indent=2)
                text = f"[ไฟล์ JSON: {name}]\n{raw_text}"
            else:
                raw_text = content.decode('utf-8', errors='ignore')
                text = f"[ไฟล์: {name}]\n{raw_text}"
        except Exception as e:
            return {"ok": False, "error": str(e)}

    extracted = auto_extract_skills(raw_text, name)
    return {"ok": True, "filename": name, "is_image": False, "text": text[:8000],
            "skills_extracted": extracted}


@router.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    content = await file.read()
    return await run_in_threadpool(
        _parse_upload, content, file.filename or "file", file.content_type or "")
